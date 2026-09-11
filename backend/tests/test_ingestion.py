"""Extracao, normalizacao e chunking.

Os casos aqui sao os que degradam a qualidade da busca em silencio: cabecalho que entra
no chunk, frase partida ao meio, `section_path` errado. Nenhum deles gera erro — apenas
piora o retrieval de um jeito que so aparece semanas depois, quando alguem reclama que
"o sistema nao acha nada".
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.modules.ingestion.chunking import dividir
from app.modules.ingestion.extractors.base import (
    Bloco,
    ExtractionError,
    NoTextLayerError,
    TextoExtraido,
    TipoBloco,
)
from app.modules.ingestion.extractors.docx import DocxExtractor
from app.modules.ingestion.extractors.pdf import PdfExtractor
from app.modules.ingestion.extractors.plaintext import MarkdownExtractor, PlainTextExtractor
from app.modules.ingestion.normalizer import normalizar
from app.modules.ingestion.tokens import contar_tokens


def docx_bytes(paragrafos: list[tuple[str, str]]) -> bytes:
    """Monta um DOCX real a partir de (estilo, texto)."""
    import docx

    documento = docx.Document()
    for estilo, texto in paragrafos:
        documento.add_paragraph(texto, style=estilo) if estilo else documento.add_paragraph(texto)
    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


# --- Extracao --------------------------------------------------------------


def test_markdown_usa_o_nivel_da_cerquilha() -> None:
    conteudo = b"# Titulo\n\nTexto do corpo.\n\n## Subtitulo\n\nMais texto."

    blocos = MarkdownExtractor().extrair(conteudo).blocos
    titulos = [(b.texto, b.nivel) for b in blocos if b.tipo is TipoBloco.TITULO]

    assert titulos == [("Titulo", 1), ("Subtitulo", 2)]


def test_markdown_remove_marcadores_de_lista() -> None:
    """`-` e `1.` nao acrescentam significado e virariam ruido no embedding."""
    blocos = MarkdownExtractor().extrair(b"- primeiro item\n- segundo item").blocos

    assert "primeiro item segundo item" in blocos[0].texto
    assert "-" not in blocos[0].texto


def test_markdown_reconhece_titulo_setext() -> None:
    blocos = MarkdownExtractor().extrair(b"Titulo Principal\n===\n\nCorpo.").blocos

    assert blocos[0].tipo is TipoBloco.TITULO
    assert blocos[0].texto == "Titulo Principal"


def test_docx_usa_o_estilo_declarado() -> None:
    """O formato AFIRMA o que e titulo; inferir por heuristica perderia precisao."""
    conteudo = docx_bytes(
        [("Heading 1", "1 Objetivo"), ("", "Texto."), ("Heading 2", "1.1 Escopo")]
    )

    blocos = DocxExtractor().extrair(conteudo).blocos
    titulos = [(b.texto, b.nivel) for b in blocos if b.tipo is TipoBloco.TITULO]

    assert titulos == [("1 Objetivo", 1), ("1.1 Escopo", 2)]


def test_texto_puro_ignora_linhas_decorativas() -> None:
    blocos = PlainTextExtractor().extrair(b"TITULO\n======\n\nCorpo do texto.").blocos

    assert all("=" not in b.texto for b in blocos)


def test_arquivo_nao_utf8_da_erro_acionavel() -> None:
    with pytest.raises(ExtractionError, match="UTF-8"):
        PlainTextExtractor().extrair("procedimento".encode("utf-16"))


def test_pdf_sem_camada_de_texto_tem_erro_proprio() -> None:
    """Erro distinto de ExtractionError: reenviar o mesmo arquivo nao resolve, o
    administrador precisa de um PDF pesquisavel."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    c.showPage()  # pagina em branco, sem texto
    c.save()

    with pytest.raises(NoTextLayerError):
        PdfExtractor().extrair(buffer.getvalue())


def test_docx_corrompido_da_erro_tratado() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("qualquer.txt", "nao e um docx")

    with pytest.raises(ExtractionError):
        DocxExtractor().extrair(buffer.getvalue())


# --- Normalizacao ----------------------------------------------------------


def test_cabecalho_repetido_e_removido() -> None:
    """Se entrar nos chunks, o nome do documento aparece em todo embedding e aproxima
    artificialmente trechos que nada tem em comum."""
    blocos = []
    for pagina in (1, 2, 3):
        blocos += [
            Bloco(texto="Banco Exemplo - Manual Interno", pagina=pagina),
            Bloco(texto=f"Pagina {pagina}", pagina=pagina),
            Bloco(texto=f"Conteudo exclusivo da pagina {pagina}.", pagina=pagina),
        ]

    resultado = normalizar(TextoExtraido(blocos=blocos, total_paginas=3))

    assert "Banco Exemplo - Manual Interno" not in resultado.texto
    assert "Pagina 1" not in resultado.texto
    assert "Conteudo exclusivo da pagina 2." in resultado.texto


def test_rodape_numerado_e_detectado_apesar_de_variar() -> None:
    """ "Pagina 1" e "Pagina 2" nao sao identicos, mas sao o mesmo rodape. Sem mascarar
    os digitos, todo rodape numerado escaparia da deteccao."""
    blocos = [Bloco(texto=f"Pagina {p} de 3", pagina=p) for p in (1, 2, 3)] + [
        Bloco(texto=f"Texto da pagina {p}.", pagina=p) for p in (1, 2, 3)
    ]

    resultado = normalizar(TextoExtraido(blocos=blocos, total_paginas=3))

    assert "Pagina 1 de 3" not in resultado.texto


def test_conteudo_repetido_no_meio_nao_e_removido() -> None:
    """Repeticao so indica cabecalho nas BORDAS da pagina. No meio, e conteudo."""
    blocos = []
    for pagina in (1, 2, 3):
        blocos += [Bloco(texto=f"Abertura {pagina}", pagina=pagina)]
        blocos += [Bloco(texto=f"Enchimento {pagina}.{i}", pagina=pagina) for i in range(8)]
        blocos += [Bloco(texto="Definicao repetida de propósito.", pagina=pagina)]
        blocos += [Bloco(texto=f"Enchimento final {pagina}.{i}", pagina=pagina) for i in range(8)]

    resultado = normalizar(TextoExtraido(blocos=blocos, total_paginas=3))

    assert "Definicao repetida de propósito." in resultado.texto


def test_frase_quebrada_pelo_pdf_e_rejuntada() -> None:
    """O PDF quebra por largura visual. Sem rejuntar, "R$" fica separado do valor e a
    busca literal por "R$ 5.000,00" falha."""
    blocos = [
        Bloco(texto="Transferencias acima de R$", pagina=1),
        Bloco(texto="5.000,00 por dia ficam bloqueadas.", pagina=1),
    ]

    resultado = normalizar(TextoExtraido(blocos=blocos, total_paginas=1))

    assert "R$ 5.000,00 por dia" in resultado.texto


def test_numeracao_de_secao_nao_e_confundida_com_continuacao() -> None:
    """ "3.2 Procedimento" apos uma linha sem ponto final e titulo novo, nao continuacao."""
    blocos = [
        Bloco(texto="O prazo vale para todos os canais", pagina=1),
        Bloco(texto="3.2 Procedimento de atualizacao", pagina=1),
    ]

    resultado = normalizar(TextoExtraido(blocos=blocos, total_paginas=1))

    assert any(s.numero == "3.2" for s in resultado.secoes)


def test_hifenizacao_de_quebra_e_desfeita() -> None:
    resultado = normalizar(TextoExtraido(blocos=[Bloco(texto="atualiza-\ncao cadastral")]))

    assert "atualizacao cadastral" in resultado.texto


def test_section_path_acumula_a_hierarquia() -> None:
    """E o caminho que aparece na citacao. Sem ele, a fonte exibida seria so o nome do
    documento — verdadeira, mas inutil para conferir o trecho."""
    blocos = [
        Bloco(texto="3 Atualizacao cadastral"),
        Bloco(texto="Texto da secao tres com conteudo suficiente."),
        Bloco(texto="3.2 Procedimento de atualizacao"),
        Bloco(texto="Texto da subsecao com conteudo suficiente."),
        Bloco(texto="4 Restricoes"),
        Bloco(texto="Texto da secao quatro."),
    ]

    secoes = normalizar(TextoExtraido(blocos=blocos)).secoes
    caminhos = {s.numero: s.caminho for s in secoes}

    assert caminhos["3"] == "3 Atualizacao cadastral"
    assert caminhos["3.2"] == "3 Atualizacao cadastral > 3.2 Procedimento de atualizacao"
    # A secao 4 fecha a 3: nao pode herdar o caminho dela.
    assert caminhos["4"] == "4 Restricoes"


def test_offsets_apontam_para_o_texto_normalizado() -> None:
    """Os offsets sao gravados no chunk e usados para destacar o trecho na UI. Se
    apontassem para o texto bruto, o destaque cairia no lugar errado."""
    blocos = [Bloco(texto="1 Objetivo"), Bloco(texto="Conteudo da secao um.")]

    resultado = normalizar(TextoExtraido(blocos=blocos))
    secao = resultado.secoes[0]

    assert resultado.texto[secao.inicio :].startswith("1 Objetivo")


# --- Chunking --------------------------------------------------------------


def _documento(texto_por_secao: dict[str, str]):
    blocos = []
    for titulo, corpo in texto_por_secao.items():
        blocos += [Bloco(texto=titulo), Bloco(texto=corpo)]
    return normalizar(TextoExtraido(blocos=blocos))


def test_chunk_nunca_atravessa_fronteira_de_secao() -> None:
    """A regra que sustenta a citacao: um chunk que comeca na 3.2 e termina na 3.3
    teria section_path de meia-verdade."""
    documento = _documento(
        {
            "1 Primeira": "Conteudo da primeira secao. " * 12,
            "2 Segunda": "Conteudo da segunda secao. " * 12,
        }
    )

    chunks = dividir(documento)

    for chunk in chunks:
        assert not ("primeira" in chunk.texto.lower() and "segunda" in chunk.texto.lower())


def test_secao_grande_e_dividida_com_sobreposicao() -> None:
    documento = _documento({"1 Longa": "Frase de conteudo relevante. " * 200})

    chunks = dividir(documento)

    assert len(chunks) > 1
    assert all(c.tokens <= 800 for c in chunks)
    # Todos herdam o caminho da mesma secao.
    assert len({c.section_path for c in chunks}) == 1


def test_secoes_curtas_sao_agrupadas() -> None:
    """Um chunk de 20 tokens casa com quase qualquer consulta sem responder nada.

    O criterio e o AGRUPAMENTO, nao um tamanho minimo absoluto: se o documento inteiro
    tem 30 tokens, o chunk tera 30 — nao ha de onde tirar mais.
    """
    documento = _documento({f"{i} Item": "Texto curto." for i in range(1, 6)})

    chunks = dividir(documento)

    assert len(chunks) == 1, f"cinco secoes curtas deveriam virar um chunk, viraram {len(chunks)}"
    assert "1 Item" in chunks[0].texto
    assert "5 Item" in chunks[0].texto


def test_offsets_do_chunk_recuperam_o_texto_exato() -> None:
    documento = _documento({"1 Secao": "Conteudo relevante para o teste. " * 20})

    for chunk in dividir(documento):
        assert documento.texto[chunk.inicio : chunk.fim] == chunk.texto


def test_documento_vazio_nao_gera_chunk() -> None:
    assert dividir(normalizar(TextoExtraido(blocos=[]))) == []


def test_documento_sem_secao_ainda_gera_chunks() -> None:
    """Texto corrido sem numeracao nem titulo: a janela deslizante cuida do resto."""
    documento = normalizar(TextoExtraido(blocos=[Bloco(texto="frase sem estrutura. " * 60)]))

    chunks = dividir(documento)

    assert chunks
    assert all(c.section_path is None for c in chunks)


# --- Tokens ----------------------------------------------------------------


def test_contagem_de_tokens_e_conservadora() -> None:
    """Superestimar produz chunks menores, que e inofensivo. Subestimar faz o contexto
    do RAG estourar o limite do modelo em producao."""
    texto = "a" * 100

    assert contar_tokens(texto) >= 100 / 4


def test_texto_vazio_tem_zero_tokens() -> None:
    assert contar_tokens("") == 0
