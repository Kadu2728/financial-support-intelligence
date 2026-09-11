"""Extracao de DOCX.

Ao contrario do PDF, o DOCX declara a estrutura: paragrafos com estilo `Heading N` sao
titulos, com nivel conhecido. Essa informacao e aproveitada diretamente — inferir por
heuristica algo que o formato ja afirma seria perder precisao de graca.

DOCX nao tem paginas: a paginacao e calculada pelo renderizador, nao armazenada. Por
isso `pagina` fica None, e o `page_number` do chunk tambem.
"""

from __future__ import annotations

import re
from io import BytesIO

from app.modules.ingestion.extractors.base import (
    Bloco,
    ExtractionError,
    TextoExtraido,
    TipoBloco,
)

_HEADING = re.compile(r"^Heading (\d)$", re.I)


class DocxExtractor:
    def extrair(self, conteudo: bytes) -> TextoExtraido:
        import docx
        from docx.opc.exceptions import PackageNotFoundError

        try:
            documento = docx.Document(BytesIO(conteudo))
        except (PackageNotFoundError, KeyError, ValueError, OSError) as exc:
            raise ExtractionError(f"DOCX ilegivel: {exc}") from exc

        blocos: list[Bloco] = []

        for paragrafo in documento.paragraphs:
            texto = paragrafo.text.strip()
            if not texto:
                continue

            estilo = (paragrafo.style.name if paragrafo.style else "") or ""
            if achado := _HEADING.match(estilo):
                blocos.append(Bloco(texto=texto, tipo=TipoBloco.TITULO, nivel=int(achado.group(1))))
            elif estilo.lower() == "title":
                blocos.append(Bloco(texto=texto, tipo=TipoBloco.TITULO, nivel=1))
            else:
                blocos.append(Bloco(texto=texto))

        # Tabelas vem depois dos paragrafos, nao intercaladas: recuperar a ordem real
        # do XML exigiria descer ao elemento bruto. Nenhum documento do acervo tem
        # tabela; quando houver, a ordem passara a importar.
        for tabela in documento.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
                if celulas:
                    blocos.append(Bloco(texto=" | ".join(celulas)))

        metadados: dict[str, str] = {}
        propriedades = documento.core_properties
        if propriedades.title:
            metadados["title"] = propriedades.title
        if propriedades.author:
            metadados["author"] = propriedades.author

        return TextoExtraido(blocos=blocos, metadados=metadados)
